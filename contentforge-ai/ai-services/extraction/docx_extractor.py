import io
import logging
from typing import Any, Dict, List, Union
import docx

logger = logging.getLogger(__name__)


def extract_text_from_docx(file_input: Union[bytes, str, io.BytesIO]) -> List[Dict[str, Any]]:
    """
    Extract text paragraph-by-paragraph and from tables in a DOCX document.

    Args:
        file_input: Raw DOCX bytes, file path string, or BytesIO buffer.

    Returns:
        List of dicts: [{"paragraph_index": int, "section": str, "text": str}]
    """
    elements: List[Dict[str, Any]] = []

    try:
        if isinstance(file_input, bytes):
            doc = docx.Document(io.BytesIO(file_input))
        elif isinstance(file_input, io.BytesIO):
            doc = docx.Document(file_input)
        elif isinstance(file_input, str):
            doc = docx.Document(file_input)
        else:
            raise ValueError(f"Unsupported file_input type for DOCX extraction: {type(file_input)}")

        current_heading = "Main"
        para_idx = 1

        # Extract paragraphs
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            # Check if this paragraph is a heading to track section
            if para.style.name.startswith("Heading"):
                current_heading = text

            elements.append({
                "paragraph_index": para_idx,
                "section": current_heading,
                "text": text,
            })
            para_idx += 1

        # Extract table text if present
        table_idx = 1
        for table in doc.tables:
            table_rows = []
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    table_rows.append(row_text)

            if table_rows:
                elements.append({
                    "paragraph_index": para_idx,
                    "section": f"Table {table_idx}",
                    "text": "\n".join(table_rows),
                })
                para_idx += 1
                table_idx += 1

        logger.info(f"Successfully extracted {len(elements)} paragraphs/elements from DOCX")
        return elements

    except Exception as e:
        logger.error(f"Error extracting text from DOCX: {e}")
        raise
