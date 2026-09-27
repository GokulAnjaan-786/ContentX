import io
import os
import sys
import uuid
from typing import Generator
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
import pymupdf
import docx

# Ensure paths
TESTS_DIR = os.path.dirname(os.path.abspath(__file__))
BACKEND_DIR = os.path.dirname(TESTS_DIR)
ROOT_DIR = os.path.dirname(BACKEND_DIR)
for p in [ROOT_DIR, BACKEND_DIR]:
    if p not in sys.path:
        sys.path.insert(0, p)

# Set test environment flags
os.environ["ENVIRONMENT"] = "testing"
os.environ["SECRET_KEY"] = "super-secret-testing-key-for-contentforge-ai-jwt"
os.environ["CLAMAV_ENABLED"] = "true"
os.environ["DATABASE_URL"] = "sqlite:///:memory:"

from app.core.config import settings
from app.core.database import Base, get_db
from app.core.security import get_password_hash, create_access_token
from app.models.organisation import Organisation
from app.models.user import User, UserRole
from app.models.document import Document
from app.models.document_chunk import DocumentChunk
from app.models.audit_log import AuditLog
from app.main import app

import app.core.database as app_db

# In-memory SQLite engine for fast, isolated tests
TEST_DATABASE_URL = "sqlite:///:memory:"
test_engine = create_engine(
    TEST_DATABASE_URL,
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=test_engine)

# Wire test_engine and sessionmaker directly into application core database
app_db.engine = test_engine
app_db.SessionLocal = TestingSessionLocal


@pytest.fixture(autouse=True)
def setup_test_db():
    """Create fresh tables for every test."""
    Base.metadata.create_all(bind=test_engine)
    yield
    Base.metadata.drop_all(bind=test_engine)


@pytest.fixture
def db_session() -> Generator:
    """Yield a database session from TestingSessionLocal."""
    session = TestingSessionLocal()
    yield session
    session.close()


@pytest.fixture
def client(db_session) -> Generator:
    """FastAPI TestClient with overridden get_db dependency."""
    def override_get_db():
        try:
            yield db_session
        finally:
            pass

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


@pytest.fixture
def test_organisation(db_session) -> Organisation:
    """Fixture to provide a test organisation."""
    org = Organisation(id=uuid.uuid4(), name="Acme Content Labs")
    db_session.add(org)
    db_session.commit()
    db_session.refresh(org)
    return org


@pytest.fixture
def test_user(db_session, test_organisation) -> User:
    """Fixture to provide a standard operator user."""
    user = User(
        id=uuid.uuid4(),
        org_id=test_organisation.id,
        email="operator@acmelabs.com",
        password_hash=get_password_hash("ValidPass123!"),
        role=UserRole.OPERATOR,
    )
    db_session.add(user)
    db_session.commit()
    db_session.refresh(user)
    return user


@pytest.fixture
def test_admin_user(db_session, test_organisation) -> User:
    """Fixture to provide an org_admin user."""
    admin = User(
        id=uuid.uuid4(),
        org_id=test_organisation.id,
        email="admin@acmelabs.com",
        password_hash=get_password_hash("AdminPass123!"),
        role=UserRole.ORG_ADMIN,
    )
    db_session.add(admin)
    db_session.commit()
    db_session.refresh(admin)
    return admin


@pytest.fixture
def auth_headers(test_user) -> dict:
    """Authorization header with valid JWT token for test_user."""
    token = create_access_token(
        subject=str(test_user.id),
        claims={"role": test_user.role.value, "org_id": str(test_user.org_id)},
    )
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def admin_auth_headers(test_admin_user) -> dict:
    """Authorization header with valid JWT token for admin user."""
    token = create_access_token(
        subject=str(test_admin_user.id),
        claims={"role": test_admin_user.role.value, "org_id": str(test_admin_user.org_id)},
    )
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def sample_pdf_bytes() -> bytes:
    """Generate a genuine 2-page PDF document in-memory using PyMuPDF."""
    doc = pymupdf.open()
    
    # Page 1
    p1 = doc.new_page()
    p1.insert_text(
        (50, 72),
        "ContentForge AI Executive Briefing\n\n"
        "Antigravity systems represent a quantum leap in autonomous orchestration. "
        "Every downstream artifact maintains strict adherence to the unified fact registry. "
        "The architecture decouples extraction from synthesis to guarantee consistency.\n\n"
        "Page 1 - Confidential"
    )
    
    # Page 2
    p2 = doc.new_page()
    p2.insert_text(
        (50, 72),
        "ContentForge AI Executive Briefing\n\n"
        "Section 2: Performance and Verification.\n"
        "Document ingestion utilizes PyMuPDF and python-docx with background processing. "
        "All uploaded documents undergo strict cryptographic signature validation.\n\n"
        "Page 2 - Confidential"
    )
    
    pdf_data = doc.write()
    doc.close()
    return pdf_data


@pytest.fixture
def sample_docx_bytes() -> bytes:
    """Generate a genuine DOCX document in-memory using python-docx."""
    doc = docx.Document()
    doc.add_heading("ContentForge AI Advisory Report", level=1)
    doc.add_paragraph("This advisory specifies requirements for multi-format content generation.")
    doc.add_paragraph("All facts are anchored to precise source sentence indexes in the registry.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def fake_pdf_bytes() -> bytes:
    """Generate a disguised executable file with DOS/PE binary signature."""
    return b"MZ\x90\x00\x03\x00\x00\x00\x04\x00\x00\x00\xff\xff\x00\x00\xb8\x00\x00\x00" * 20


@pytest.fixture
def eicar_virus_bytes() -> bytes:
    """Standard EICAR antivirus test file content."""
    return b"X5O!P%@AP[4\\PZX54(P^)7CC)7}$EICAR-STANDARD-ANTIVIRUS-TEST-FILE!$H+H*"
